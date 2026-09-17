import os
import re
import sys
import json
import struct
import bisect
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side


SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
EXPORT_DIR = os.path.join(SCRIPT_DIR, 'export')

INSTR_FUNC_MAP = {}
INSTR_COLOR_MAP = {}
MAP_FILE = os.path.join(SCRIPT_DIR, 'evt_instr_map.json')
if os.path.exists(MAP_FILE):
    with open(MAP_FILE, 'r', encoding='utf-8') as f:
        raw_map = json.load(f)
    for k, v in raw_map.items():
        if isinstance(v, dict):
            INSTR_FUNC_MAP[k] = v.get('func', '')
            if v.get('color'):
                INSTR_COLOR_MAP[k] = v['color']
        else:
            INSTR_FUNC_MAP[k] = v


# ── 公共工具 ──────────────────────────────────────────

def xor_decrypt(data, key=0xFF):
    return bytes([b ^ key for b in data])


def read_u32_le(data, offset):
    return struct.unpack_from('<I', data, offset)[0]


def clean_text(text):
    return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)


def decode_big5(data):
    try:
        text = data.decode('big5', errors='replace')
        text = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', text)
        text = text.replace('\r', '').replace('\n', '')
        return text
    except Exception:
        return ''


def get_instr_prefix(hex_str):
    parts = hex_str.split(' ')
    if len(parts) >= 5:
        b0 = parts[0]
        b4 = parts[4]
        return '{}000000{}'.format(b0, b4)
    return ''


# ── 样式 ──────────────────────────────────────────────

header_font = Font(name='微软雅黑', bold=True, color='FFFFFF', size=11)
header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
cell_align = Alignment(vertical='top', wrap_text=True)
thin_border = Border(
    left=Side(style='thin'),
    right=Side(style='thin'),
    top=Side(style='thin'),
    bottom=Side(style='thin')
)
skip_fill = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
seg_fill = PatternFill(start_color='D9E2F3', end_color='D9E2F3', fill_type='solid')
last_fill = PatternFill(start_color='FFF2CC', end_color='FFF2CC', fill_type='solid')


def apply_header(ws, headers):
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border


def set_col_widths(ws, widths):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + i)].width = w


# ── EVT 解析 ──────────────────────────────────────────

def parse_evt(decrypted):
    entry_count = decrypted[10]
    menu_start = 274

    all_entries = []
    for i in range(entry_count):
        off = menu_start + i * 16
        if off + 16 > len(decrypted):
            break
        seg_id = read_u32_le(decrypted, off)
        count = read_u32_le(decrypted, off + 4)
        record_size = read_u32_le(decrypted, off + 8)
        total_size = read_u32_le(decrypted, off + 12)
        all_entries.append({
            'offset': off, 'id': seg_id, 'count': count,
            'record_size': record_size, 'total_size': total_size,
        })

    all_seg_ids = sorted(set(e['id'] for e in all_entries if e['id'] != 0))

    valid_entries = []
    seen_ids = set()
    for entry in all_entries:
        if entry['count'] == 0:
            continue
        if entry['id'] in seen_ids:
            continue
        seen_ids.add(entry['id'])
        valid_entries.append(entry)

    all_instructions = []
    for seg in valid_entries:
        seg_id = seg['id']
        count = seg['count']

        idx_in_sorted = bisect.bisect_left(all_seg_ids, seg_id)
        if idx_in_sorted + 1 < len(all_seg_ids):
            next_seg_start = all_seg_ids[idx_in_sorted + 1]
        else:
            next_seg_start = len(decrypted)

        if seg_id == 0x0442:
            record_size = seg.get('record_size', 0)
            total_size = seg.get('total_size', 0)
            if record_size == 0 and total_size > 0 and count > 0:
                record_size = total_size // count
            if count == 1 and record_size == 0:
                record_size = total_size

            menu_offset = seg_id
            for i in range(count):
                rec_off = seg_id + i * record_size
                if i + 1 < count:
                    end_off = seg_id + (i + 1) * record_size
                else:
                    end_off = min(seg_id + total_size, next_seg_start) if total_size > 0 else next_seg_start

                if rec_off >= len(decrypted):
                    inst_hex = '[超出文件范围]'
                    inst_len = 0
                else:
                    actual_end = min(end_off, len(decrypted))
                    inst_bytes = decrypted[rec_off:actual_end]
                    if all(b == 0 for b in inst_bytes):
                        continue
                    inst_hex = ' '.join('{:02X}'.format(b) for b in inst_bytes)
                    inst_len = len(inst_bytes)

                all_instructions.append({
                    'seg_id': '0x{:04X}'.format(seg_id),
                    'seg_offset': '0x{:04X}'.format(menu_offset),
                    'inst_index': i,
                    'inst_offset': '0x{:04X}'.format(rec_off),
                    'end_offset': '0x{:04X}'.format(end_off),
                    'length': inst_len,
                    'hex': inst_hex,
                    'is_first_seg': True,
                })
        else:
            menu_offset = seg_id
            entries = []
            for i in range(count):
                entry_off = menu_offset + i * 4
                if entry_off + 4 > len(decrypted):
                    break
                inst_off = read_u32_le(decrypted, entry_off)
                entries.append(inst_off)

            for i, inst_off in enumerate(entries):
                if i + 1 < len(entries):
                    end_off = entries[i + 1]
                else:
                    end_off = next_seg_start

                if inst_off >= len(decrypted):
                    inst_hex = '[超出文件范围]'
                    inst_len = 0
                else:
                    actual_end = min(end_off, len(decrypted))
                    inst_bytes = decrypted[inst_off:actual_end]
                    inst_hex = ' '.join('{:02X}'.format(b) for b in inst_bytes)
                    inst_len = len(inst_bytes)

                all_instructions.append({
                    'seg_id': '0x{:04X}'.format(seg_id),
                    'seg_offset': '0x{:04X}'.format(menu_offset),
                    'inst_index': i,
                    'inst_offset': '0x{:04X}'.format(inst_off),
                    'end_offset': '0x{:04X}'.format(end_off),
                    'length': inst_len,
                    'hex': inst_hex,
                    'is_first_seg': False,
                })

    return all_entries, valid_entries, all_instructions


# ── MSG 解析 ──────────────────────────────────────────

def parse_msg(decrypted):
    entry_count = decrypted[10]
    menu_start = 274

    all_entries = []
    for i in range(entry_count):
        off = menu_start + i * 16
        if off + 16 > len(decrypted):
            break
        seg_id = read_u32_le(decrypted, off)
        count = read_u32_le(decrypted, off + 4)
        all_entries.append({'offset': off, 'id': seg_id, 'count': count})

    all_seg_ids = sorted(set(e['id'] for e in all_entries if e['id'] != 0))

    valid_entries = []
    seen_ids = set()
    for entry in all_entries:
        if entry['count'] == 0:
            continue
        if entry['id'] in seen_ids:
            continue
        seen_ids.add(entry['id'])
        valid_entries.append(entry)

    all_instructions = []
    for seg in valid_entries:
        seg_id = seg['id']
        count = seg['count']
        menu_offset = seg_id

        raw_entries = []
        for i in range(count + 1):
            entry_off = menu_offset + i * 4
            if entry_off + 4 > len(decrypted):
                break
            inst_off = read_u32_le(decrypted, entry_off)
            raw_entries.append(inst_off)

        valid_raw = [e for e in raw_entries if e != 0]
        if not valid_raw:
            continue

        idx_in_sorted = bisect.bisect_left(all_seg_ids, seg_id)
        if idx_in_sorted + 1 < len(all_seg_ids):
            next_seg_start = all_seg_ids[idx_in_sorted + 1]
        else:
            next_seg_start = len(decrypted)

        for i, inst_off in enumerate(valid_raw):
            if i + 1 < len(valid_raw):
                end_off = valid_raw[i + 1]
            else:
                end_off = next_seg_start

            if inst_off >= len(decrypted):
                inst_hex = '[超出文件范围]'
                inst_text = ''
                inst_len = 0
            else:
                actual_end = min(end_off, len(decrypted))
                inst_bytes = decrypted[inst_off:actual_end]
                inst_hex = ' '.join('{:02X}'.format(b) for b in inst_bytes)
                inst_len = len(inst_bytes)
                inst_text = clean_text(inst_bytes.decode('big5', errors='replace'))

            all_instructions.append({
                'seg_id': '0x{:04X}'.format(seg_id),
                'seg_offset': '0x{:04X}'.format(menu_offset),
                'inst_index': i,
                'inst_offset': '0x{:04X}'.format(inst_off),
                'end_offset': '0x{:04X}'.format(end_off),
                'length': inst_len,
                'hex': inst_hex,
                'text': inst_text,
            })

    return all_entries, valid_entries, all_instructions


# ── DAT 解析 ──────────────────────────────────────────

def parse_dat(decrypted):
    entry_count = decrypted[10]
    menu_start = 274

    all_entries = []
    for i in range(entry_count):
        off = menu_start + i * 16
        if off + 16 > len(decrypted):
            break
        seg_id = read_u32_le(decrypted, off)
        count = read_u32_le(decrypted, off + 4)
        record_size = read_u32_le(decrypted, off + 8)
        total_size = read_u32_le(decrypted, off + 12)
        all_entries.append({
            'index': i,
            'offset': off,
            'seg_id': seg_id,
            'count': count,
            'record_size': record_size,
            'total_size': total_size,
        })

    all_seg_ids = sorted(set(e['seg_id'] for e in all_entries if e['seg_id'] != 0))

    valid_entries = []
    seen_ids = set()
    for entry in all_entries:
        if entry['count'] == 0:
            continue
        if entry['seg_id'] == 0:
            continue
        if entry['seg_id'] in seen_ids:
            continue
        seen_ids.add(entry['seg_id'])
        valid_entries.append(entry)

    all_records = []
    for entry in valid_entries:
        seg_id = entry['seg_id']
        count = entry['count']
        record_size = entry['record_size']

        if record_size == 0:
            if entry['total_size'] > 0 and count > 0:
                record_size = entry['total_size'] // count
            else:
                record_size = 0

        if count == 1 and record_size == 0:
            record_size = entry['total_size']

        idx_in_sorted = bisect.bisect_left(all_seg_ids, seg_id)
        if idx_in_sorted + 1 < len(all_seg_ids):
            next_seg_start = all_seg_ids[idx_in_sorted + 1]
        else:
            next_seg_start = len(decrypted)

        for i in range(count):
            rec_off = seg_id + i * record_size
            if i + 1 < count:
                end_off = seg_id + (i + 1) * record_size
            else:
                end_off = min(seg_id + entry['total_size'], next_seg_start) if entry['total_size'] > 0 else next_seg_start

            if rec_off >= len(decrypted):
                rec_hex = '[超出文件范围]'
                rec_len = 0
                rec_big5 = ''
                rec_group = ''
                rec_num = ''
                rec_name = ''
                rec_coord = ''
                rec_remark = ''
                rec_bytes = b''
            else:
                actual_end = min(end_off, len(decrypted))
                rec_bytes = decrypted[rec_off:actual_end]
                rec_hex = ' '.join('{:02X}'.format(b) for b in rec_bytes)
                if len(rec_bytes) >= 4 and rec_bytes[3] == 0x00 and not all(b == 0 for b in rec_bytes) and (rec_bytes[0] != 0 or rec_bytes[1] != 0 or rec_bytes[2] != 0):
                    rec_group = '{:02X}'.format(rec_bytes[0])
                    rec_num = '{:02X}{:02X}'.format(rec_bytes[1], rec_bytes[2])
                    is_char_data = True
                else:
                    rec_group = ''
                    rec_num = ''
                    is_char_data = False
                if is_char_data and len(rec_bytes) >= 12:
                    name_end = rec_bytes.find(b'\x00', 4)
                    if name_end == -1 or name_end == 4:
                        name_bytes = rec_bytes[4:12].rstrip(b'\x00')
                    else:
                        name_bytes = rec_bytes[4:name_end]
                    rec_name = clean_text(name_bytes.decode('big5', errors='replace')) if name_bytes else ''
                elif not is_char_data:
                    if all(b == 0 for b in rec_bytes):
                        rec_group = None
                        rec_num = None
                        rec_name = None
                        rec_coord = None
                    else:
                        name_raw = rec_bytes.rstrip(b'\x00').decode('big5', errors='replace')
                        rec_name = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', name_raw)
                else:
                    rec_name = ''
                if is_char_data and len(rec_bytes) >= 28:
                    rec_x = struct.unpack_from('<H', rec_bytes, 24)[0]
                    rec_y = struct.unpack_from('<H', rec_bytes, 26)[0]
                    rec_coord = '{}, {}'.format(rec_x, rec_y)
                else:
                    rec_coord = ''
                if is_char_data and len(rec_bytes) >= 30 and rec_bytes[29] != 0xFF:
                    rec_remark = '击败后跳转第{}指令组'.format(rec_bytes[29])
                else:
                    rec_remark = ''
                rec_len = len(rec_bytes)
                rec_big5 = decode_big5(rec_bytes)

            if all(b == 0 for b in rec_bytes):
                continue

            all_records.append({
                'seg_index': entry['index'],
                'seg_id': '0x{:04X}'.format(seg_id),
                'seg_offset': '0x{:04X}'.format(seg_id),
                'record_index': i,
                'record_offset': '0x{:04X}'.format(rec_off),
                'end_offset': '0x{:04X}'.format(end_off),
                'length': rec_len,
                'hex': rec_hex,
                'big5': rec_big5,
                'group': rec_group,
                'num': rec_num,
                'name': rec_name,
                'coord': rec_coord,
                'remark': rec_remark,
            })

    return all_entries, valid_entries, all_records


# ── SWL 解析（补充文件，定长记录结构）────────────────────

def parse_swl(decrypted):
    if len(decrypted) < 3 or decrypted[0:3] != b'SE3':
        return None, None, None
    entry_count = decrypted[10]
    menu_start = 274

    all_entries = []
    for i in range(entry_count):
        off = menu_start + i * 16
        if off + 16 > len(decrypted):
            break
        seg_id = read_u32_le(decrypted, off)
        count = read_u32_le(decrypted, off + 4)
        record_size = read_u32_le(decrypted, off + 8)
        total_size = read_u32_le(decrypted, off + 12)
        all_entries.append({
            'offset': off, 'id': seg_id, 'count': count,
            'record_size': record_size, 'total_size': total_size,
        })

    all_seg_ids = sorted(set(e['id'] for e in all_entries if e['id'] != 0))

    valid_entries = []
    seen_ids = set()
    for entry in all_entries:
        if entry['count'] == 0:
            continue
        if entry['id'] in seen_ids:
            continue
        seen_ids.add(entry['id'])
        valid_entries.append(entry)

    all_instructions = []
    for seg in valid_entries:
        seg_id = seg['id']
        count = seg['count']

        idx_in_sorted = bisect.bisect_left(all_seg_ids, seg_id)
        if idx_in_sorted + 1 < len(all_seg_ids):
            next_seg_start = all_seg_ids[idx_in_sorted + 1]
        else:
            next_seg_start = len(decrypted)

        record_size = seg.get('record_size', 0)
        total_size = seg.get('total_size', 0)
        if record_size == 0 and total_size > 0 and count > 0:
            record_size = total_size // count
        if count == 1 and record_size == 0:
            record_size = total_size

        menu_offset = seg_id
        for i in range(count):
            rec_off = seg_id + i * record_size
            if i + 1 < count:
                end_off = seg_id + (i + 1) * record_size
            else:
                end_off = min(seg_id + total_size, next_seg_start) if total_size > 0 else next_seg_start

            if rec_off >= len(decrypted):
                inst_hex = '[超出文件范围]'
                inst_len = 0
            else:
                actual_end = min(end_off, len(decrypted))
                inst_bytes = decrypted[rec_off:actual_end]
                if all(b == 0 for b in inst_bytes):
                    continue
                inst_hex = ' '.join('{:02X}'.format(b) for b in inst_bytes)
                inst_len = len(inst_bytes)

            all_instructions.append({
                'seg_id': '0x{:04X}'.format(seg_id),
                'seg_offset': '0x{:04X}'.format(menu_offset),
                'inst_index': i,
                'inst_offset': '0x{:04X}'.format(rec_off),
                'end_offset': '0x{:04X}'.format(end_off),
                'length': inst_len,
                'hex': inst_hex,
            })

    return all_entries, valid_entries, all_instructions


# ── 主程序 ─────────────────────────────────────────────

def process_file(file_path, ext):
    with open(file_path, 'rb') as f:
        raw = f.read()

    decrypted = xor_decrypt(raw)

    if len(decrypted) < 3 or decrypted[0:3] != b'SE3':
        print('  [{}] {} -> 文件格式不正确（非SE3格式），跳过'.format(ext.upper(), os.path.basename(file_path)))
        return None, None, None

    if ext == '.evt':
        return parse_evt(decrypted)
    elif ext == '.msg':
        return parse_msg(decrypted)
    elif ext == '.dat':
        return parse_dat(decrypted)
    elif ext == '.swl':
        return parse_swl(decrypted)
    return None, None, None


def save_stage_excel(stage_name, evt_data, msg_data, dat_data, swl_data, output_path):
    wb = Workbook()

    # ── Sheet 1: 整合表 ──
    ws1 = wb.active
    ws1.title = '整合表'
    headers1 = ['二级菜单Offset', '指令序号', '指令内容(HEX)', '具体指令', '指令作用', '指令内容']
    apply_header(ws1, headers1)
    set_col_widths(ws1, [18, 10, 40, 14, 16, 30])
    ws1.freeze_panes = 'A2'

    msg_lookup = []
    if msg_data and msg_data[2]:
        seg_map = {}
        for mi in msg_data[2]:
            so = mi['seg_offset']
            if so not in seg_map:
                seg_map[so] = []
            seg_map[so].append(mi)
        msg_lookup = [seg_map[k] for k in sorted(seg_map.keys())]

    dat_group_names = {}
    if dat_data and dat_data[2]:
        for rec in dat_data[2]:
            g = rec.get('group', '')
            n = rec.get('name', '')
            c = rec.get('coord', '')
            if g and n:
                dat_group_names.setdefault(g, []).append((n, c))

    current_seg = None
    row_idx = 0
    if evt_data and evt_data[2]:
        for inst in evt_data[2]:
            row_idx += 1
            is_new_seg = inst['seg_offset'] != current_seg
            current_seg = inst['seg_offset']
            is_first_seg = inst.get('is_first_seg', False)
            if is_first_seg:
                prefix = ''
                func = ''
                content = ''
                if inst['hex'] != '[超出文件范围]':
                    hex_bytes = inst['hex'].split(' ')
                    if len(hex_bytes) >= 31:
                        byte22 = int(hex_bytes[22], 16)
                        x = struct.unpack_from('<H', bytes([int(hex_bytes[23], 16), int(hex_bytes[24], 16)]), 0)[0]
                        y = struct.unpack_from('<H', bytes([int(hex_bytes[25], 16), int(hex_bytes[26], 16)]), 0)[0]
                        byte29 = int(hex_bytes[29], 16)
                        byte30 = int(hex_bytes[30], 16)

                        if byte22 == 0x00:
                            if byte29 == 0x01:
                                type_str = '关卡切换坐标'
                            elif byte29 == 0x00:
                                type_str = '战场事件坐标'
                            else:
                                type_str = '未知坐标'
                        elif byte22 == 0x01:
                            if byte30 != 0xFF:
                                type_str = '特殊事件坐标'
                            else:
                                if byte29 == 0x00:
                                    type_str = '战场隐藏道具坐标'
                                else:
                                    type_str = '普通可见道具坐标'
                        else:
                            type_str = '未知坐标'

                        content = '{}（{}，{}）'.format(type_str, x, y)
                        if byte30 != 0xFF:
                            seg_id_str = ''
                            if evt_data and evt_data[0] and byte30 < len(evt_data[0]):
                                seg_id_str = '0x{:04X}'.format(evt_data[0][byte30]['id'])
                            content += '，跳转到{}指令组（{}）'.format(byte30, seg_id_str)
            else:
                prefix = get_instr_prefix(inst['hex']) if inst['hex'] != '[超出文件范围]' else ''
                func = INSTR_FUNC_MAP.get(prefix, '') if prefix else ''
                content = ''
            if not is_first_seg and inst['hex'] != '[超出文件范围]':
                hex_bytes = inst['hex'].split(' ')
                if prefix in ('1400000014', '1500000015'):
                    if len(hex_bytes) >= 12:
                        seg_idx = int(hex_bytes[10], 16)
                        dlg_idx = int(hex_bytes[11], 16)
                        if seg_idx < len(msg_lookup) and dlg_idx < len(msg_lookup[seg_idx]):
                            content = '"{}"'.format(msg_lookup[seg_idx][dlg_idx].get('text', ''))
                elif prefix == '0E0000000E':
                    if len(hex_bytes) >= 17:
                        jump_val = struct.unpack_from('<I', bytes([int(h, 16) for h in hex_bytes[13:17]]), 0)[0]
                        content = '跳转至{}指令'.format(jump_val)
                elif prefix == '0600000006':
                    name_bytes = []
                    for h in hex_bytes[6:]:
                        b = int(h, 16)
                        if b == 0:
                            break
                        name_bytes.append(b)
                    name = clean_text(bytes(name_bytes).decode('big5', errors='replace')) if name_bytes else ''
                    content = name
                elif prefix == '0C0000000C':
                    if len(hex_bytes) > 5:
                        start_idx = 5
                        if int(hex_bytes[5], 16) == 0 and len(hex_bytes) > 6:
                            start_idx = 6
                        name_bytes = []
                        for h in hex_bytes[start_idx:]:
                            b = int(h, 16)
                            if b == 0:
                                break
                            name_bytes.append(b)
                        if name_bytes:
                            content = clean_text(bytes(name_bytes).decode('big5', errors='replace'))
                elif prefix == '0100000001':
                    if len(hex_bytes) > 5:
                        grp = '{:02X}'.format(int(hex_bytes[5], 16))
                        if grp in dat_group_names:
                            parts = []
                            for item in dat_group_names[grp]:
                                name = item[0]
                                coord = item[1]
                                if coord:
                                    parts.append('{}（{}）'.format(name, coord))
                                else:
                                    parts.append(name)
                            content = '\n'.join(parts)
                elif prefix == '1600000016':
                    if len(hex_bytes) > 5:
                        grp = '{:02X}'.format(int(hex_bytes[5], 16))
                        if grp in dat_group_names:
                            names = [item[0] for item in dat_group_names[grp]]
                            content = '\n'.join(names)
                elif prefix == '1000000010':
                    if len(hex_bytes) >= 12:
                        jump_val = int(hex_bytes[11], 16)
                        content = '跳转至{}指令'.format(jump_val)
                elif prefix == '1E0000001E':
                    if len(hex_bytes) > 5:
                        name_bytes = []
                        for h in hex_bytes[5:]:
                            b = int(h, 16)
                            if b == 0:
                                break
                            name_bytes.append(b)
                        if name_bytes:
                            content = clean_text(bytes(name_bytes).decode('big5', errors='replace'))
                elif prefix == '1F0000001F':
                    if len(hex_bytes) >= 8:
                        mode = int(hex_bytes[5], 16)
                        music_id = int(hex_bytes[6], 16)
                        action = '开始播放' if mode == 0 else '停止播放'
                        content = '{}音乐{}.mp3'.format(action, music_id)
                elif prefix == '1200000012':
                    if len(hex_bytes) >= 10:
                        x = struct.unpack_from('<H', bytes([int(hex_bytes[6], 16), int(hex_bytes[7], 16)]), 0)[0]
                        y = struct.unpack_from('<H', bytes([int(hex_bytes[8], 16), int(hex_bytes[9], 16)]), 0)[0]
                        content = '（{}，{}）'.format(x, y)
                elif prefix == '1300000013':
                    if len(hex_bytes) >= 30:
                        x = struct.unpack_from('<H', bytes([int(hex_bytes[26], 16), int(hex_bytes[27], 16)]), 0)[0]
                        y = struct.unpack_from('<H', bytes([int(hex_bytes[28], 16), int(hex_bytes[29], 16)]), 0)[0]
                        name_start = 6 if int(hex_bytes[5], 16) == 0 else 5
                        name_bytes = []
                        for h in hex_bytes[name_start:]:
                            b = int(h, 16)
                            if b == 0:
                                break
                            name_bytes.append(b)
                        name = clean_text(bytes(name_bytes).decode('big5', errors='replace')) if name_bytes else ''
                        content = '{}移动到坐标 ({}, {})'.format(name, x, y)
                elif prefix == '1700000017':
                    if len(hex_bytes) >= 30:
                        x = struct.unpack_from('<H', bytes([int(hex_bytes[26], 16), int(hex_bytes[27], 16)]), 0)[0]
                        y = struct.unpack_from('<H', bytes([int(hex_bytes[28], 16), int(hex_bytes[29], 16)]), 0)[0]
                        name_start = 6 if int(hex_bytes[5], 16) == 0 else 5
                        name_bytes = []
                        for h in hex_bytes[name_start:]:
                            b = int(h, 16)
                            if b == 0:
                                break
                            name_bytes.append(b)
                        name = clean_text(bytes(name_bytes).decode('big5', errors='replace')) if name_bytes else ''
                        content = '{}移动到坐标 ({}, {})'.format(name, x, y)
                elif prefix == '1D0000001D':
                    pass
                elif prefix == '0A0000000A':
                    if len(hex_bytes) > 5:
                        mode = hex_bytes[5]
                        if mode == '00':
                            content = '切换到战斗模式'
                        elif mode == '01':
                            content = '切换到RPG模式'
                elif prefix == '1B0000001B':
                    if len(hex_bytes) > 6:
                        ascii_bytes = []
                        for h in hex_bytes[6:]:
                            b = int(h, 16)
                            if b == 0:
                                break
                            ascii_bytes.append(b)
                        filename = bytes(ascii_bytes).decode('ascii', errors='replace') if ascii_bytes else ''
                        big5_start = 6 + len(ascii_bytes) + 1
                        while big5_start < len(hex_bytes) and int(hex_bytes[big5_start], 16) == 0:
                            big5_start += 1
                        big5_bytes = []
                        for h in hex_bytes[big5_start:]:
                            b = int(h, 16)
                            if b == 0:
                                break
                            big5_bytes.append(b)
                        big5_text = clean_text(bytes(big5_bytes).decode('big5', errors='replace')) if big5_bytes else ''
                        if filename and big5_text:
                            content = '{}（{}）'.format(big5_text, filename)
                        elif filename:
                            content = filename
                        elif big5_text:
                            content = big5_text
                elif prefix == '2B0000002B':
                    if len(hex_bytes) >= 9:
                        x = int(hex_bytes[5], 16)
                        y = int(hex_bytes[7], 16)
                        content = '坐标（{}，{}）'.format(x, y)
                elif prefix == '2E0000002E':
                    if len(hex_bytes) >= 30:
                        name_start = 6 if int(hex_bytes[5], 16) == 0 else 5
                        name_bytes = []
                        for h in hex_bytes[name_start:]:
                            b = int(h, 16)
                            if b == 0:
                                break
                            name_bytes.append(b)
                        name = clean_text(bytes(name_bytes).decode('big5', errors='replace')) if name_bytes else ''
                        x = struct.unpack_from('<H', bytes([int(hex_bytes[26], 16), int(hex_bytes[27], 16)]), 0)[0]
                        y = struct.unpack_from('<H', bytes([int(hex_bytes[28], 16), int(hex_bytes[29], 16)]), 0)[0]
                        content = '{}（{}，{}）'.format(name, x, y)
                elif prefix == '0D0000000D':
                    if len(hex_bytes) >= 26:
                        name_bytes = []
                        for h in hex_bytes[5:]:
                            b = int(h, 16)
                            if b == 0:
                                break
                            name_bytes.append(b)
                        level_name = clean_text(bytes(name_bytes).decode('big5', errors='replace'))
                        jump_target = int(hex_bytes[25], 16)
                        content = '跳转到{}关卡执行第{}组指令'.format(level_name, jump_target)
                elif prefix == '2200000022':
                    name_bytes = []
                    for h in hex_bytes[6:]:
                        b = int(h, 16)
                        if b == 0:
                            break
                        name_bytes.append(b)
                    name = clean_text(bytes(name_bytes).decode('big5', errors='replace')) if name_bytes else ''
                    content = name
                elif prefix == '3A0000003A':
                    name_bytes = []
                    for h in hex_bytes[5:]:
                        b = int(h, 16)
                        if b == 0:
                            break
                        name_bytes.append(b)
                    name = clean_text(bytes(name_bytes).decode('big5', errors='replace')) if name_bytes else ''
                    content = name
                elif prefix == '3900000039':
                    name_bytes = []
                    for h in hex_bytes[6:]:
                        b = int(h, 16)
                        if b == 0:
                            break
                        name_bytes.append(b)
                    name = clean_text(bytes(name_bytes).decode('big5', errors='replace')) if name_bytes else ''
                    content = name
                elif prefix == '2000000020':
                    if len(hex_bytes) >= 11:
                        b5 = hex_bytes[5]
                        jump_target = int(hex_bytes[10], 16)
                        if b5 == '01':
                            content = '倒计时停止'
                        else:
                            seg_id_str = ''
                            if evt_data and evt_data[0] and jump_target < len(evt_data[0]):
                                seg_id_str = '0x{:04X}'.format(evt_data[0][jump_target]['id'])
                            mode = '显式开始' if b5 == '00' else '隐式开始'
                            content = '{}，倒计时结束触发{}指令组（{}）'.format(mode, jump_target, seg_id_str)
            values = [
                inst['seg_offset'] if is_new_seg else '',
                inst['inst_index'],
                inst['hex'],
                prefix,
                func,
                content,
            ]
            row_fill = None
            if not is_first_seg and prefix and prefix in INSTR_COLOR_MAP:
                row_fill = PatternFill(start_color=INSTR_COLOR_MAP[prefix],
                                       end_color=INSTR_COLOR_MAP[prefix],
                                       fill_type='solid')
            for col, v in enumerate(values, 1):
                if isinstance(v, str):
                    v = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]', '', v)
                cell = ws1.cell(row=row_idx + 1, column=col, value=v)
                cell.alignment = cell_align
                cell.border = thin_border
                if row_fill:
                    cell.fill = row_fill

    # ── Sheet 2: 指令菜单 ──
    if evt_data and evt_data[0] is not None:
        ws2 = wb.create_sheet('指令菜单')
        all_entries = evt_data[0]
        valid_ids = {e['id'] for e in evt_data[1]} if evt_data[1] else set()
        headers2 = ['序号', '一级指令段ID', '一级菜单Offset', '指令数量', '是否跳过', '作用', '备注']
        apply_header(ws2, headers2)

        SEG_ROLES = {
            0: '绑定坐标事件和道具',
            1: '主脚本',
            5: '战斗结束后脚本',
            6: 'GAME OVER事件',
            11: '每击杀一名敌人调用一次',
        }

        for idx, entry in enumerate(all_entries):
            seg_id = entry['id']
            count = entry['count']
            is_skipped = (count == 0) or (seg_id not in valid_ids)
            note = ''
            if count == 0:
                note = '无指令(count=0)'
            elif seg_id not in valid_ids:
                note = '重复段ID'
            elif seg_id == 0x0442:
                note = '第一段（定长记录）'

            role = SEG_ROLES.get(idx, '')

            values = [
                idx,
                '0x{:04X}'.format(seg_id),
                '0x{:04X}'.format(entry['offset']),
                count,
                '是' if is_skipped else '否',
                role,
                note,
            ]
            for col, v in enumerate(values, 1):
                cell = ws2.cell(row=idx + 2, column=col, value=v)
                cell.alignment = header_align
                cell.border = thin_border
                if is_skipped:
                    cell.fill = skip_fill

        set_col_widths(ws2, [6, 18, 18, 10, 10, 18, 20])
        ws2.freeze_panes = 'A2'

    # ── Sheet 3: 指令解析 ──
    if evt_data and evt_data[2] is not None:
        ws3 = wb.create_sheet('指令解析')
        all_instructions = evt_data[2]
        headers3 = ['序号', '所属一级段ID', '二级菜单Offset', '指令序号',
                    '指令Offset', '结束Offset', '长度(字节)', '指令内容(HEX)',
                    '具体指令', '指令作用']
        apply_header(ws3, headers3)

        current_seg = None
        for idx, inst in enumerate(all_instructions):
            is_new_seg = inst['seg_id'] != current_seg
            current_seg = inst['seg_id']

            is_last = False
            if idx + 1 < len(all_instructions):
                next_inst = all_instructions[idx + 1]
                if next_inst['seg_id'] != current_seg:
                    is_last = True
            else:
                is_last = True

            if inst.get('is_first_seg', False):
                prefix = ''
                func = ''
            else:
                prefix = get_instr_prefix(inst['hex']) if inst['hex'] != '[超出文件范围]' else ''
                func = INSTR_FUNC_MAP.get(prefix, '') if prefix else ''

            values = [
                idx,
                inst['seg_id'] if is_new_seg else '',
                inst['seg_offset'] if is_new_seg else '',
                inst['inst_index'],
                inst['inst_offset'],
                inst['end_offset'],
                inst['length'],
                inst['hex'],
                prefix,
                func,
            ]
            for col, v in enumerate(values, 1):
                cell = ws3.cell(row=idx + 2, column=col, value=v)
                cell.alignment = cell_align
                cell.border = thin_border
                if is_new_seg:
                    cell.fill = seg_fill
                elif is_last:
                    cell.fill = last_fill

        set_col_widths(ws3, [6, 16, 18, 10, 14, 14, 12, 80, 14, 16])
        ws3.freeze_panes = 'A2'

    # ── Sheet 4: 对话菜单 ──
    if msg_data and msg_data[0] is not None:
        ws4 = wb.create_sheet('对话菜单')
        all_entries = msg_data[0]
        valid_ids = {e['id'] for e in msg_data[1]} if msg_data[1] else set()
        headers4 = ['序号', '一级指令段ID', '一级菜单Offset', '指令数量', '是否跳过', '备注']
        apply_header(ws4, headers4)

        for idx, entry in enumerate(all_entries):
            seg_id = entry['id']
            count = entry['count']
            is_skipped = (count == 0) or (seg_id == 0) or (seg_id not in valid_ids)
            note = ''
            if count == 0:
                note = '无指令(count=0)'
            elif seg_id == 0:
                note = '空段ID'
            elif seg_id not in valid_ids:
                note = '重复段ID'

            values = [
                idx,
                '0x{:04X}'.format(seg_id),
                '0x{:04X}'.format(entry['offset']),
                count,
                '是' if is_skipped else '否',
                note,
            ]
            for col, v in enumerate(values, 1):
                cell = ws4.cell(row=idx + 2, column=col, value=v)
                cell.alignment = header_align
                cell.border = thin_border
                if is_skipped:
                    cell.fill = skip_fill

        set_col_widths(ws4, [6, 18, 18, 10, 10, 20])
        ws4.freeze_panes = 'A2'

    # ── Sheet 5: 对话解析 ──
    if msg_data and msg_data[2] is not None:
        ws5 = wb.create_sheet('对话解析')
        all_instructions = msg_data[2]
        headers5 = ['序号', '所属一级段ID', '二级菜单Offset', '指令序号',
                    '指令Offset', '结束Offset', '长度(字节)', '指令内容(HEX)', '文字内容(Big5)']
        apply_header(ws5, headers5)

        current_seg = None
        row_idx = 0
        for inst in all_instructions:
            if inst['length'] == 0:
                continue
            is_new_seg = inst['seg_id'] != current_seg
            current_seg = inst['seg_id']

            is_last = False

            values = [
                row_idx,
                inst['seg_id'] if is_new_seg else '',
                inst['seg_offset'] if is_new_seg else '',
                inst['inst_index'],
                inst['inst_offset'],
                inst['end_offset'],
                inst['length'],
                inst['hex'],
                inst['text'],
            ]
            for col, v in enumerate(values, 1):
                cell = ws5.cell(row=row_idx + 2, column=col, value=v)
                cell.alignment = cell_align
                cell.border = thin_border
                if is_new_seg:
                    cell.fill = seg_fill
            row_idx += 1

        set_col_widths(ws5, [6, 16, 18, 10, 14, 14, 12, 60, 60])
        ws5.freeze_panes = 'A2'

    # ── Sheet 6: 数据菜单 ──
    if dat_data and dat_data[0] is not None:
        ws6 = wb.create_sheet('数据菜单')
        all_entries = dat_data[0]
        valid_ids = {e['seg_id'] for e in dat_data[1]} if dat_data[1] else set()
        headers6 = ['序号', '数据段ID', '菜单Offset', '记录数', '记录大小(字节)', '总大小(字节)', '是否跳过', '备注']
        apply_header(ws6, headers6)

        for idx, entry in enumerate(all_entries):
            seg_id = entry['seg_id']
            count = entry['count']
            is_skipped = (count == 0) or (seg_id == 0) or (seg_id not in valid_ids)
            note = ''
            if count == 0:
                note = '无记录(count=0)'
            elif seg_id == 0:
                note = '段ID为0'
            elif seg_id not in valid_ids:
                note = '重复段ID'

            values = [
                idx,
                '0x{:04X}'.format(seg_id),
                '0x{:04X}'.format(entry['offset']),
                count,
                entry.get('record_size', 0),
                entry.get('total_size', 0),
                '是' if is_skipped else '否',
                note,
            ]
            for col, v in enumerate(values, 1):
                cell = ws6.cell(row=idx + 2, column=col, value=v)
                cell.alignment = header_align
                cell.border = thin_border
                if is_skipped:
                    cell.fill = skip_fill

        set_col_widths(ws6, [6, 16, 16, 10, 16, 16, 10, 20])
        ws6.freeze_panes = 'A2'

    # ── Sheet 7: 数据解析 ──
    if dat_data and dat_data[2] is not None:
        ws7 = wb.create_sheet('数据解析')
        all_records = dat_data[2]
        headers7 = ['序号', '所属数据段ID', '数据段Offset', '记录序号',
                    '记录Offset', '结束Offset', '长度(字节)',
                    '记录内容(HEX)', '分组', '编号', '姓名', '坐标(X,Y)', '备注']
        apply_header(ws7, headers7)

        current_seg = None
        for idx, rec in enumerate(all_records):
            is_new_seg = rec['seg_id'] != current_seg
            current_seg = rec['seg_id']

            values = [
                idx,
                rec['seg_id'] if is_new_seg else '',
                rec['seg_offset'] if is_new_seg else '',
                rec['record_index'],
                rec['record_offset'],
                rec['end_offset'],
                rec['length'],
                rec['hex'],
                rec['group'],
                rec['num'],
                rec['name'],
                rec['coord'],
                rec.get('remark', ''),
            ]
            for col, v in enumerate(values, 1):
                cell = ws7.cell(row=idx + 2, column=col, value=v)
                cell.alignment = cell_align
                cell.border = thin_border
                if is_new_seg:
                    cell.fill = seg_fill

        set_col_widths(ws7, [6, 16, 16, 10, 14, 14, 12, 80, 8, 10, 12, 12, 24])
        ws7.freeze_panes = 'A2'

    # ── Sheet 8: 补充菜单 ──
    if swl_data and swl_data[0] is not None:
        ws8 = wb.create_sheet('补充菜单')
        all_entries = swl_data[0]
        valid_ids = {e['id'] for e in swl_data[1]} if swl_data[1] else set()
        headers8 = ['序号', '补充指令段ID', '菜单Offset', '指令数量', '是否跳过', '备注']
        apply_header(ws8, headers8)

        for idx, entry in enumerate(all_entries):
            seg_id = entry['id']
            count = entry['count']
            is_skipped = (count == 0) or (seg_id not in valid_ids)
            note = ''
            if count == 0:
                note = '无指令(count=0)'
            elif seg_id not in valid_ids:
                note = '重复段ID'

            values = [
                idx,
                '0x{:04X}'.format(seg_id),
                '0x{:04X}'.format(entry['offset']),
                count,
                '是' if is_skipped else '否',
                note,
            ]
            for col, v in enumerate(values, 1):
                cell = ws8.cell(row=idx + 2, column=col, value=v)
                cell.alignment = header_align
                cell.border = thin_border
                if is_skipped:
                    cell.fill = skip_fill

        set_col_widths(ws8, [6, 18, 18, 10, 10, 20])
        ws8.freeze_panes = 'A2'

    # ── Sheet 9: 补充解析 ──
    if swl_data and swl_data[2] is not None:
        ws9 = wb.create_sheet('补充解析')
        all_instructions = swl_data[2]
        headers9 = ['序号', '所属补充段ID', '二级菜单Offset', '指令序号',
                    '指令Offset', '结束Offset', '长度(字节)', '指令内容(HEX)',
                    '指令内容']
        apply_header(ws9, headers9)

        current_seg = None
        for idx, inst in enumerate(all_instructions):
            is_new_seg = inst['seg_id'] != current_seg
            current_seg = inst['seg_id']

            content = ''
            if inst['hex'] != '[超出文件范围]':
                hex_bytes = inst['hex'].split(' ')
                if len(hex_bytes) >= 23:
                    byte14 = int(hex_bytes[14], 16)
                    x = struct.unpack_from('<H', bytes([int(hex_bytes[15], 16), int(hex_bytes[16], 16)]), 0)[0]
                    y = struct.unpack_from('<H', bytes([int(hex_bytes[17], 16), int(hex_bytes[18], 16)]), 0)[0]
                    byte21 = int(hex_bytes[21], 16)
                    byte22 = int(hex_bytes[22], 16)

                    if byte14 == 0x00:
                        if byte21 == 0x01:
                            type_str = '关卡切换坐标'
                        elif byte21 == 0x00:
                            type_str = '战场事件坐标'
                        else:
                            type_str = '未知坐标'
                    elif byte14 == 0x01:
                        if byte22 != 0xFF:
                            type_str = '特殊事件坐标'
                        else:
                            if byte21 == 0x00:
                                type_str = '战场隐藏道具坐标'
                            else:
                                type_str = '普通可见道具坐标'
                    else:
                        type_str = '未知坐标'

                    content = '{}（{}，{}）'.format(type_str, x, y)
                    if byte22 != 0xFF:
                        seg_id_str = ''
                        if evt_data and evt_data[0] and byte22 < len(evt_data[0]):
                            seg_id_str = '0x{:04X}'.format(evt_data[0][byte22]['id'])
                        content += '，跳转到{}指令组（{}）'.format(byte22, seg_id_str)

            values = [
                idx,
                inst['seg_id'] if is_new_seg else '',
                inst['seg_offset'] if is_new_seg else '',
                inst['inst_index'],
                inst['inst_offset'],
                inst['end_offset'],
                inst['length'],
                inst['hex'],
                content,
            ]
            for col, v in enumerate(values, 1):
                cell = ws9.cell(row=idx + 2, column=col, value=v)
                cell.alignment = cell_align
                cell.border = thin_border
                if is_new_seg:
                    cell.fill = seg_fill

        set_col_widths(ws9, [6, 16, 18, 10, 14, 14, 12, 80, 30])
        ws9.freeze_panes = 'A2'

    try:
        wb.save(output_path)
    except PermissionError:
        print('  [跳过] 文件被占用，无法写入 - {}'.format(os.path.basename(output_path)))
        return


def main():
    os.chdir(SCRIPT_DIR)
    print('=' * 50)
    print('  致命武力2脚本批量解析工具 (EVT / MSG / DAT / SWL)')
    print('=' * 50)
    print()

    if len(sys.argv) >= 2:
        folder = sys.argv[1].strip().strip('"')
    else:
        folder = input('请输入脚本文件夹路径: ').strip().strip('"')
    if not folder:
        print('未输入路径，程序退出。')
        return
    if not os.path.isdir(folder):
        print('错误: 文件夹不存在 - {}'.format(folder))
        return

    if not os.path.exists(EXPORT_DIR):
        os.makedirs(EXPORT_DIR)

    files = []
    for f in sorted(os.listdir(folder)):
        ext = os.path.splitext(f)[1].lower()
        if ext in ('.evt', '.msg', '.dat', '.swl'):
            files.append((os.path.join(folder, f), ext))

    if not files:
        print('文件夹中没有 .evt / .msg / .dat / .swl 文件。')
        return

    print('找到 {} 个文件:'.format(len(files)))
    for fp, ext in files:
        print('  {} ({})'.format(os.path.basename(fp), ext))
    print()

    parsed = {}
    for fp, ext in files:
        name = os.path.basename(fp)
        try:
            data = process_file(fp, ext)
            if data[0] is not None:
                parsed.setdefault(name, {})[ext] = data
                label = {'.evt': '指令', '.msg': '对话', '.dat': '记录', '.swl': '补充'}[ext]
                count = len(data[2])
                print('  [{}] {} -> {} 条{}'.format(ext.upper(), name, count, label))
        except Exception as e:
            print('  [{}] {} -> 错误: {}'.format(ext.upper(), name, e))

    stage_groups = {}
    for fname, types in parsed.items():
        for ftype, data in types.items():
            base = re.sub(r'^(?:[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}_)+', '', os.path.splitext(fname)[0])
            stage_groups.setdefault(base, {})[ftype] = data

    for stage_name, types in stage_groups.items():
        evt_data = types.get('.evt', (None, None, None))
        msg_data = types.get('.msg', (None, None, None))
        dat_data = types.get('.dat', (None, None, None))
        swl_data = types.get('.swl', (None, None, None))
        output_name = '{}.xlsx'.format(stage_name)
        save_stage_excel(stage_name, evt_data, msg_data, dat_data, swl_data, os.path.join(EXPORT_DIR, output_name))
        print('  [EXPORT] {} -> {}'.format(stage_name, output_name))

    print()
    print('完成! 所有 Excel 文件已保存到:')
    print('  export 文件夹: {}'.format(EXPORT_DIR))
    print()
    try:
        input('按回车键退出...')
    except EOFError:
        pass


if __name__ == '__main__':
    main()
